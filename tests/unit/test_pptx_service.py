import re
from io import BytesIO

from pptx import Presentation

from src.services.pptx_service import PptxService


def get_presentation(bytes_data):
    return Presentation(BytesIO(bytes_data))

def find_text(slide, pattern):
    regex = re.compile(pattern, re.IGNORECASE)
    for shape in slide.shapes:
        if hasattr(shape, "text") and regex.search(shape.text):
            return shape.text
    return None

def get_slide_notes(slide):
    try:
        return slide.notes_slide.notes_text_frame.text.strip()
    except Exception:
        return None

BASIC_SAMPLE = {
    "title": "My Presentation",
    "slides": [
        {"type": "title", "title": "Welcome", "subtitle": "Intro", "notes": "Say hello!"},
        {"type": "content", "title": "About", "bullets": ["A", "B", "C"], "notes": "Main content."},
        {"type": "two-column", "title": "Compare", "left": {"heading": "L", "bullets": ["1"]}, "right": {"heading": "R", "bullets": ["2"]}, "notes": "Columns..."},
        {"type": "section", "title": "Next Section", "notes": ""},
        {"type": "content", "title": "Finale", "bullets": ["Bye"], "notes": ""}
    ]
}

def test_generate_basic():
    pptx_bytes = PptxService().generate(BASIC_SAMPLE)
    assert isinstance(pptx_bytes, bytes)
    # Should be a valid pptx file
    prs = get_presentation(pptx_bytes)
    assert prs is not None
    assert len(prs.slides) == 5

def test_slide_count():
    n = 5
    d = {"slides": [
        {"type": "content", "title": f"Slide {i}", "bullets": [f"Item {i}"], "notes": str(i)} for i in range(n)
    ]}
    pptx_bytes = PptxService().generate(d)
    prs = get_presentation(pptx_bytes)
    assert len(prs.slides) == n

def test_title_slide():
    sample = {"slides": [{"type": "title", "title": "Main", "subtitle": "Sub", "notes": ""}]}
    pptx_bytes = PptxService().generate(sample)
    prs = get_presentation(pptx_bytes)
    slide = prs.slides[0]
    assert find_text(slide, r"main") is not None
    assert find_text(slide, r"sub") is not None

def test_content_slide():
    sample = {"slides": [{"type": "content", "title": "CT", "bullets": ["A", "B"], "notes": "N"}]}
    pptx_bytes = PptxService().generate(sample)
    prs = get_presentation(pptx_bytes)
    slide = prs.slides[0]
    assert find_text(slide, r"ct")
    all_text = " ".join(s.text for s in slide.shapes if hasattr(s, "text"))
    assert "A" in all_text and "B" in all_text

def test_two_column_slide():
    sample = {"slides": [{"type": "two-column", "title": "ColTest", "left": {"heading": "H1", "bullets": ["a1"]}, "right": {"heading": "H2", "bullets": ["b2"]}, "notes": ""}]}
    pptx_bytes = PptxService().generate(sample)
    prs = get_presentation(pptx_bytes)
    slide = prs.slides[0]
    all_texts = " ".join([s.text for s in slide.shapes if hasattr(s, "text")])
    assert "H1" in all_texts
    assert "H2" in all_texts
    assert "a1" in all_texts
    assert "b2" in all_texts

def test_section_slide():
    sample = {"slides": [{"type": "section", "title": "Section Break", "notes": "N"}]}
    pptx_bytes = PptxService().generate(sample)
    prs = get_presentation(pptx_bytes)
    slide = prs.slides[0]
    assert find_text(slide, r"section break")

def test_speaker_notes():
    pptx_bytes = PptxService().generate(BASIC_SAMPLE)
    prs = get_presentation(pptx_bytes)
    notes = [get_slide_notes(sl) for sl in prs.slides]
    assert "Say hello!" in notes[0]
    assert "Main content." in notes[1]
    assert "Columns..." in notes[2]

def test_watermark():
    pptx_bytes = PptxService().generate(BASIC_SAMPLE, watermark=True)
    prs = get_presentation(pptx_bytes)
    found = False
    for slide in prs.slides:
        for shape in slide.shapes:
            if hasattr(shape, "text") and "slide" in shape.text.lower() and "made" in shape.text.lower():
                found = True
    assert found

def test_no_watermark():
    pptx_bytes = PptxService().generate(BASIC_SAMPLE, watermark=False)
    prs = get_presentation(pptx_bytes)
    count = 0
    for slide in prs.slides:
        for shape in slide.shapes:
            if hasattr(shape, "text") and "slide" in shape.text.lower() and "made" in shape.text.lower():
                count += 1
    assert count == 0

def test_education_theme():
    pptx_bytes = PptxService().generate(BASIC_SAMPLE, theme="education")
    prs = get_presentation(pptx_bytes)
    assert len(prs.slides) == 5

def test_empty_slides():
    sample = {"slides": [], "title": "Blank"}
    pptx_bytes = PptxService().generate(sample)
    prs = get_presentation(pptx_bytes)
    assert isinstance(pptx_bytes, bytes)
    # Should be valid pptx, probably 1 slide (title) if title present, or 0 otherwise
    assert len(prs.slides) in (0, 1)
